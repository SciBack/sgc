"""Importador de la estructura institucional desde hoja de cálculo (#48).

Una instancia recién instalada tiene DocTypes, roles y workflows, pero no sabe nada de la
institución, y casi todo el modelo cuelga de su estructura: los procesos tienen unidad
propietaria, los indicadores se miden por programa y periodo. Esto carga esa estructura desde
las hojas en que las instituciones la tienen hoy:

    unidades        codigo*, nombre*, tipo, padre, sede, es_campus
    programas       codigo*, nombre*, codigo_oficial, nivel, facultad, escuela, estado
    programas-sede  codigo*, programa*, sede*, escuela, licencia, estado
    periodos        codigo*, anio, semestre, fecha_inicio, fecha_fin, estado

(* obligatoria). `padre`, `sede`, `facultad`, `escuela` y `programa` son CÓDIGOS, de la misma
hoja o ya existentes en el sitio. Las fechas van como AAAA-MM-DD. El periodo vigente es el que
está `abierto`.

La fuente es una carpeta con `<hoja>.csv` (UTF-8) o un `.xlsx` con hojas de esos nombres; las
hojas que falten simplemente no se importan. Plantillas con un ejemplo, que se importan tal
cual: `sgc/public/plantillas/estructura/` (en el sitio, `/assets/sgc/plantillas/estructura/`).

**O entra todo o no entra nada.** Primero se valida el conjunto entero contra lo que ya hay en
el sitio (obligatorios, códigos duplicados, padres que no existen, ciclos, más de una raíz,
valores fuera de la lista, fechas al revés) y, si algo falla, no se escribe ninguna fila y se
dice cuál y por qué. Después se escribe todo en una sola transacción: un error a mitad la
deshace entera. Una estructura organizativa a medias es peor que ninguna.

Reimportar la misma hoja no duplica: lo existente se actualiza por código, y solo si cambia.
No borra nada que no esté en la hoja.

Ejecutar (sin `aplicar` solo simula: valida y dice qué haría):
    bench --site <sitio> execute sgc.setup.importador_estructura.run --kwargs "{'ruta': '/ruta/estructura'}"
    bench --site <sitio> execute sgc.setup.importador_estructura.run --kwargs "{'ruta': '/ruta/estructura.xlsx', 'aplicar': 1}"
"""

import csv
import os

import frappe
from frappe.utils import getdate

# Orden de escritura: cada hoja puede referirse a las anteriores.
HOJAS = ("unidades", "programas", "programas-sede", "periodos")

ESQUEMA = {
    "unidades": {
        "doctype": "Unidad Organica",
        "campos": {"codigo": "codigo", "nombre": "nombre", "tipo": "tipo", "padre": "parent_unidad_organica",
                   "sede": "sede", "es_campus": "es_campus"},
        "obligatorias": ("codigo", "nombre"),
    },
    "programas": {
        "doctype": "Programa",
        "campos": {"codigo": "codigo", "nombre": "nombre", "codigo_oficial": "codigo_oficial", "nivel": "nivel",
                   "facultad": "facultad", "escuela": "escuela", "estado": "estado"},
        "obligatorias": ("codigo", "nombre"),
    },
    "programas-sede": {
        "doctype": "Programa Sede",
        "campos": {"codigo": "codigo", "programa": "programa", "sede": "sede", "escuela": "escuela",
                   "licencia": "licencia", "estado": "estado"},
        "obligatorias": ("codigo", "programa", "sede"),
    },
    "periodos": {
        "doctype": "Periodo Academico",
        "campos": {"codigo": "codigo", "anio": "anio", "semestre": "semestre", "fecha_inicio": "fecha_inicio",
                   "fecha_fin": "fecha_fin", "estado": "estado"},
        "obligatorias": ("codigo",),
    },
}

PLANTILLAS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                          "public", "plantillas", "estructura")


class EstructuraInvalida(Exception):
    """La hoja no se puede importar. `errores` dice dónde: «hoja, fila N: …»."""

    def __init__(self, errores):
        self.errores = errores
        super().__init__(f"{len(errores)} error(es); no se ha importado nada:\n  " + "\n  ".join(errores))


# ---------------------------------------------------------------- lectura


def _limpia(fila):
    return {str(k).strip().lower(): ("" if v is None else str(v).strip()) for k, v in fila.items() if k is not None}


def leer(ruta):
    """{hoja: [filas]} desde una carpeta de CSV o un .xlsx. Cada fila lleva `_fila` (su línea en la hoja)."""
    if not os.path.exists(ruta):
        raise EstructuraInvalida([f"no existe «{ruta}»"])
    hojas = {}
    if os.path.isdir(ruta):
        for hoja in HOJAS:
            fichero = os.path.join(ruta, f"{hoja}.csv")
            if os.path.isfile(fichero):
                with open(fichero, encoding="utf-8-sig", newline="") as fh:
                    hojas[hoja] = [{**_limpia(f), "_fila": n} for n, f in enumerate(csv.DictReader(fh), start=2)]
    elif ruta.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        libro = load_workbook(ruta, read_only=True, data_only=True)
        for hoja in HOJAS:
            if hoja not in libro.sheetnames:
                continue
            filas = list(libro[hoja].iter_rows(values_only=True))
            if not filas:
                continue
            cabecera = [str(c or "").strip().lower() for c in filas[0]]
            hojas[hoja] = [
                {**_limpia(dict(zip(cabecera, _celdas(f), strict=False))), "_fila": n}
                for n, f in enumerate(filas[1:], start=2) if any(c not in (None, "") for c in f)
            ]
    else:
        raise EstructuraInvalida([f"«{ruta}» no es una carpeta de CSV ni un .xlsx"])
    for hoja, filas in hojas.items():
        hojas[hoja] = [f for f in filas if any(v for k, v in f.items() if k != "_fila")]
    return hojas


def _celdas(fila):
    """Las fechas de Excel llegan como datetime: se pasan a AAAA-MM-DD."""
    return [c.strftime("%Y-%m-%d") if hasattr(c, "strftime") else c for c in fila]


# ---------------------------------------------------------------- validación (pura)


def ordenar_por_profundidad(unidades, padre_existente=None):
    """Las unidades en un orden en que el padre siempre va antes que el hijo.

    `padre_existente` (codigo -> padre) es el árbol ya cargado en el sitio, para las unidades que
    la hoja no trae. Lanza ValueError si hay un ciclo.
    """
    padre_existente = padre_existente or {}
    padre = {u["codigo"]: u.get("padre") or None for u in unidades}
    profundidad = {}

    def _prof(codigo, visto):
        if codigo in profundidad:
            return profundidad[codigo]
        if codigo in visto:
            raise ValueError(f"ciclo en la jerarquía: {' → '.join([*visto, codigo])}")
        p = padre[codigo] if codigo in padre else padre_existente.get(codigo)
        d = 0 if not p or (p not in padre and p not in padre_existente) else _prof(p, [*visto, codigo]) + 1
        profundidad[codigo] = d
        return d

    for codigo in padre:
        _prof(codigo, [])
    return sorted(unidades, key=lambda u: (profundidad[u["codigo"]], u["_fila"]))


def validar(hojas, sitio):
    """Lista de errores del conjunto, contra lo que ya hay en el sitio. [] si se puede importar.

    `sitio` describe lo existente: {"unidades": {codigo: {"padre", "tipo"}}, "programas": set,
    "opciones": {(hoja, columna): [valores válidos]}}. Pura, para poder probarla sin base.
    """
    errores = []

    def err(hoja, fila, texto):
        errores.append(f"{hoja}, fila {fila['_fila']}: {texto}")

    # Obligatorias, duplicados y valores de lista, en todas las hojas.
    for hoja, filas in hojas.items():
        esquema = ESQUEMA[hoja]
        desconocidas = sorted({k for f in filas for k in f if k != "_fila" and k not in esquema["campos"]})
        if desconocidas:
            errores.append(f"{hoja}: columnas desconocidas {', '.join(desconocidas)} "
                           f"(válidas: {', '.join(esquema['campos'])})")
        vistos = {}
        for f in filas:
            for col in esquema["obligatorias"]:
                if not f.get(col):
                    err(hoja, f, f"falta «{col}»")
            codigo = f.get("codigo")
            if codigo:
                if codigo in vistos:
                    err(hoja, f, f"código «{codigo}» repetido (ya en la fila {vistos[codigo]})")
                vistos.setdefault(codigo, f["_fila"])
            for col, valor in f.items():
                validos = sitio["opciones"].get((hoja, col))
                if valor and validos is not None and valor not in validos:
                    err(hoja, f, f"«{col}» = «{valor}» no es válido (admite: {', '.join(v for v in validos if v)})")

    # Unidades: padres, sedes, ciclos y una sola raíz.
    unidades = hojas.get("unidades", [])
    tipo = {c: u.get("tipo") for c, u in sitio["unidades"].items()}
    tipo.update({u["codigo"]: u.get("tipo") for u in unidades if u.get("codigo")})
    for u in unidades:
        if u.get("padre") and u["padre"] not in tipo:
            err("unidades", u, f"el padre «{u['padre']}» no existe ni en la hoja ni en el sitio")
        if u.get("padre") and u["padre"] == u.get("codigo"):
            err("unidades", u, "una unidad no puede ser su propio padre")
        if u.get("sede") and tipo.get(u["sede"]) != "Sede":
            err("unidades", u, f"la sede «{u['sede']}» no es una unidad de tipo Sede")
    if unidades and not errores:
        existentes = {c: d.get("padre") for c, d in sitio["unidades"].items()}
        try:
            ordenar_por_profundidad([u for u in unidades if u.get("codigo")], existentes)
        except ValueError as e:
            errores.append(f"unidades: {e}")
        padre_final = {**existentes, **{u["codigo"]: u.get("padre") or None for u in unidades}}
        raices = sorted(c for c, p in padre_final.items() if not p)
        if len(raices) > 1:
            errores.append(f"unidades: el árbol resultante tendría {len(raices)} raíces ({', '.join(raices[:10])}"
                           f"{'…' if len(raices) > 10 else ''}); debe tener una sola, la institución")

    # Programas y programas por sede: sus referencias.
    def _unidad(hoja, f, col, tipo_esperado=None):
        valor = f.get(col)
        if not valor:
            return
        if valor not in tipo:
            err(hoja, f, f"«{col}» = «{valor}» no es una unidad de la hoja ni del sitio")
        elif tipo_esperado and tipo.get(valor) != tipo_esperado:
            err(hoja, f, f"«{col}» = «{valor}» no es una unidad de tipo {tipo_esperado}")

    for p in hojas.get("programas", []):
        _unidad("programas", p, "facultad")
        _unidad("programas", p, "escuela")
    programas = set(sitio["programas"]) | {p["codigo"] for p in hojas.get("programas", []) if p.get("codigo")}
    for ps in hojas.get("programas-sede", []):
        if ps.get("programa") and ps["programa"] not in programas:
            err("programas-sede", ps, f"el programa «{ps['programa']}» no existe ni en la hoja ni en el sitio")
        _unidad("programas-sede", ps, "sede", "Sede")
        _unidad("programas-sede", ps, "escuela")

    # Periodos: año y fechas.
    for p in hojas.get("periodos", []):
        if p.get("anio") and not p["anio"].isdigit():
            err("periodos", p, f"«anio» = «{p['anio']}» no es un año")
        fechas = {}
        for col in ("fecha_inicio", "fecha_fin"):
            if p.get(col):
                try:
                    fechas[col] = getdate(p[col])
                except Exception:
                    err("periodos", p, f"«{col}» = «{p[col]}» no es una fecha AAAA-MM-DD")
        if len(fechas) == 2 and fechas["fecha_fin"] < fechas["fecha_inicio"]:
            err("periodos", p, "la fecha de fin es anterior a la de inicio")

    return errores


# ---------------------------------------------------------------- sitio y escritura


def _opciones():
    opciones = {}
    for hoja, esquema in ESQUEMA.items():
        meta = frappe.get_meta(esquema["doctype"])
        for col, campo in esquema["campos"].items():
            df = meta.get_field(campo)
            if df and df.fieldtype == "Select":
                opciones[(hoja, col)] = (df.options or "").split("\n")
            elif df and df.fieldtype == "Check":
                opciones[(hoja, col)] = ["0", "1"]
    return opciones


def describir_sitio():
    return {
        "unidades": {u.name: {"padre": u.parent_unidad_organica, "tipo": u.tipo}
                     for u in frappe.get_all("Unidad Organica", fields=["name", "parent_unidad_organica", "tipo"])},
        "programas": set(frappe.get_all("Programa", pluck="name")),
        "opciones": _opciones(),
    }


def _valores(hoja, fila):
    campos = ESQUEMA[hoja]["campos"]
    valores = {}
    for col, campo in campos.items():
        if col not in fila:
            continue  # columna ausente: no se toca lo que haya
        v = fila[col] or None
        if campo == "es_campus":
            v = int(v or 0)
        elif v is not None and campo == "anio":
            v = int(v)
        elif v is not None and campo == "codigo_oficial":
            v = v.upper()  # `Programa.validate` lo normaliza así: compararlo igual
        valores[campo] = v
    return valores


def _igual(actual, nuevo):
    """Mismo valor para la base: vacío, None y 0 son lo mismo; una fecha se compara como texto."""
    if not actual and not nuevo:
        return True
    return str(actual if actual is not None else "") == str(nuevo if nuevo is not None else "")


def _escribir(hoja, fila, extra=None):
    """Crea o actualiza la fila por código. Devuelve «creados», «actualizados» o «sin cambios»."""
    try:
        return _escribir_fila(hoja, fila, extra)
    except Exception as e:
        raise EstructuraInvalida([f"{hoja}, fila {fila['_fila']}: {e}"]) from e


def _escribir_fila(hoja, fila, extra):
    doctype = ESQUEMA[hoja]["doctype"]
    valores = {**_valores(hoja, fila), **(extra or {})}
    codigo = valores["codigo"]
    if frappe.db.exists(doctype, codigo):
        doc = frappe.get_doc(doctype, codigo)
        cambios = {k: v for k, v in valores.items() if k != "codigo" and not _igual(doc.get(k), v)}
        if not cambios:
            return "sin cambios"
        doc.update(cambios)
        doc.save(ignore_permissions=True)
        return "actualizados"
    doc = frappe.get_doc({"doctype": doctype, **{k: v for k, v in valores.items() if v is not None}})
    doc.insert(ignore_permissions=True)
    return "creados"


def _escribir_unidades(unidades, existentes, resultado):
    """El árbol en dos pasadas: primero unidades y padres (por profundidad), después las sedes.

    La sede de una unidad es otra unidad, y puede venir más abajo en la hoja o a la misma
    profundidad: escribirla en la primera pasada fallaría por un enlace que aún no existe.
    `is_group` lo exige NestedSet en toda unidad con hijos; se deriva del árbol final.
    """
    padre_final = {**existentes, **{u["codigo"]: u.get("padre") or None for u in unidades}}
    con_hijos = {p for p in padre_final.values() if p}
    estado = {}
    for u in unidades:
        sin_sede = {k: v for k, v in u.items() if k != "sede"}
        estado[u["codigo"]] = _escribir("unidades", sin_sede, {"is_group": 1 if u["codigo"] in con_hijos else 0})
    for u in unidades:
        if "sede" in u and _escribir("unidades", {"_fila": u["_fila"], "codigo": u["codigo"], "sede": u["sede"]}) \
                == "actualizados" and estado[u["codigo"]] == "sin cambios":
            estado[u["codigo"]] = "actualizados"
    for valor in estado.values():
        resultado["unidades"][valor] += 1


def comprobar_arbol():
    """Errores de coherencia del árbol de unidades (lft/rgt y raíz única). [] si está bien."""
    filas = frappe.get_all("Unidad Organica", fields=["name", "lft", "rgt", "parent_unidad_organica"])
    por_nombre = {f.name: f for f in filas}
    errores = []
    raices = [f.name for f in filas if not f.parent_unidad_organica]
    if filas and len(raices) != 1:
        errores.append(f"el árbol tiene {len(raices)} raíces")
    for f in filas:
        if not (f.lft and f.rgt and f.lft < f.rgt):
            errores.append(f"{f.name}: lft/rgt incoherentes ({f.lft}, {f.rgt})")
        padre = por_nombre.get(f.parent_unidad_organica)
        if padre and not (padre.lft < f.lft and f.rgt < padre.rgt):
            errores.append(f"{f.name}: no está dentro de su padre {padre.name}")
    return errores


def importar(ruta, aplicar=False):
    """Valida y, con `aplicar`, escribe. Devuelve {hoja: {creados, actualizados, sin cambios}}.

    Lanza EstructuraInvalida si algo no cuadra, antes de escribir nada. Con `aplicar`, o se
    confirma todo o se deshace todo.
    """
    hojas = leer(ruta)
    if not hojas:
        raise EstructuraInvalida([f"«{ruta}» no trae ninguna hoja ({', '.join(HOJAS)})"])
    sitio = describir_sitio()
    errores = validar(hojas, sitio)
    if errores:
        raise EstructuraInvalida(errores)

    resultado = {hoja: {"creados": 0, "actualizados": 0, "sin cambios": 0} for hoja in hojas}
    existentes = {c: d["padre"] for c, d in sitio["unidades"].items()}
    if "unidades" in hojas:
        hojas["unidades"] = ordenar_por_profundidad(hojas["unidades"], existentes)

    frappe.db.savepoint("importador_estructura")
    try:
        for hoja in HOJAS:
            if hoja not in hojas:
                continue
            try:
                if hoja == "unidades":
                    _escribir_unidades(hojas[hoja], existentes, resultado)
                    continue
                for fila in hojas[hoja]:
                    resultado[hoja][_escribir(hoja, fila)] += 1
            except EstructuraInvalida:
                raise
            except Exception as e:
                raise EstructuraInvalida([f"{hoja}: {e}"]) from e
        if "unidades" in hojas:
            arbol = comprobar_arbol()
            if arbol:
                raise EstructuraInvalida([f"unidades: {e}" for e in arbol])
    except Exception:
        frappe.db.rollback(save_point="importador_estructura")
        raise
    if not aplicar:
        # Simulación: se ha escrito y comprobado de verdad, y se deshace.
        frappe.db.rollback(save_point="importador_estructura")
    return resultado


def run(ruta=None, aplicar=0):
    """Comando: importa la estructura de `ruta`, o la simula si no se pasa `aplicar`."""
    if not ruta:
        print(f"Falta `ruta`. Plantillas de ejemplo en {PLANTILLAS}")
        raise SystemExit(2)
    try:
        resultado = importar(ruta, aplicar=bool(int(aplicar)))
    except EstructuraInvalida as e:
        print(f"Estructura NO importada — {e}")
        raise SystemExit(1) from None
    print(("ESTRUCTURA IMPORTADA" if int(aplicar) else "SIMULACIÓN (sin `aplicar`: no se ha guardado nada)") + f" — {ruta}")
    for hoja, c in resultado.items():
        print(f"  {hoja}: {c['creados']} creados · {c['actualizados']} actualizados · {c['sin cambios']} sin cambios")
    if int(aplicar):
        frappe.db.commit()
